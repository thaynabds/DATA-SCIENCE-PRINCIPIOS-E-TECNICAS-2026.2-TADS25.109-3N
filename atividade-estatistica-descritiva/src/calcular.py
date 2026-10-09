"""Calcula as medidas da gasolina comum e gera a planilha e os graficos."""

from collections import Counter
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

import matplotlib.pyplot as plt
import pandas as pd
from matplotlib.ticker import FuncFormatter
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
RAIZ = Path(__file__).resolve().parents[1]
ARQUIVO_CSV = RAIZ / "dados" / "Dados_Gasolina_ANP_2026-09-21_a_27.csv"
PASTA_ENTREGA = RAIZ / "entrega"

AZUL = "1F4E79"
CINZA = "F4F6F8"
LINHA = "BFCBD9"


def ler_precos():
    tabela = pd.read_csv(ARQUIVO_CSV, sep=";", encoding="utf-8-sig")
    precos = [Decimal(str(valor).replace(",", ".")) for valor in tabela["preco_venda_rs_litro"]]
    return tabela, precos


def mediana_de(valores):
    ordem = sorted(valores)
    quantidade = len(ordem)
    meio = quantidade // 2
    if quantidade % 2 == 1:
        return ordem[meio]
    return (ordem[meio - 1] + ordem[meio]) / 2


def calcular(precos):
    quantidade = len(precos)
    ordem = sorted(precos)
    soma = sum(precos, Decimal("0"))
    media = soma / quantidade
    mediana = ordem[quantidade // 2]
    metade_inferior = ordem[: quantidade // 2]
    metade_superior = ordem[quantidade // 2 + 1 :]
    quartil_1 = mediana_de(metade_inferior)
    quartil_3 = mediana_de(metade_superior)
    intervalo = quartil_3 - quartil_1
    limite_inferior = quartil_1 - (Decimal("1.5") * intervalo)
    limite_superior = quartil_3 + (Decimal("1.5") * intervalo)
    frequencia = Counter(precos)
    maior_frequencia = max(frequencia.values())
    moda = sorted(valor for valor, vezes in frequencia.items() if vezes == maior_frequencia)
    soma_quadrados = sum((preco - media) ** 2 for preco in precos)
    variancia = soma_quadrados / (quantidade - 1)
    desvio = variancia.sqrt()
    return {
        "quantidade": quantidade,
        "ordem": ordem,
        "soma": soma,
        "media": media,
        "mediana": mediana,
        "moda": moda[0],
        "frequencia_moda": maior_frequencia,
        "frequencia": frequencia,
        "minimo": ordem[0],
        "maximo": ordem[-1],
        "amplitude": ordem[-1] - ordem[0],
        "quartil_1": quartil_1,
        "quartil_3": quartil_3,
        "intervalo": intervalo,
        "limite_inferior": limite_inferior,
        "limite_superior": limite_superior,
        "soma_quadrados": soma_quadrados,
        "variancia": variancia,
        "desvio": desvio,
    }


def dinheiro(valor, casas="0.00"):
    return valor.quantize(Decimal(casas), rounding=ROUND_HALF_UP)


def conferir(resultado):
    esperado = {
        "quantidade": 51,
        "soma": Decimal("355.21"),
        "media": Decimal("355.21") / Decimal(51),
        "mediana": Decimal("6.97"),
        "moda": Decimal("6.95"),
        "frequencia_moda": 21,
        "minimo": Decimal("6.89"),
        "maximo": Decimal("6.99"),
        "amplitude": Decimal("0.10"),
        "quartil_1": Decimal("6.95"),
        "quartil_3": Decimal("6.99"),
        "intervalo": Decimal("0.04"),
    }
    for campo, valor in esperado.items():
        if resultado[campo] != valor:
            raise SystemExit(f"Conferencia falhou em {campo}: {resultado[campo]} != {valor}")
    desvio_texto = f"{resultado['desvio']:.12f}"
    if not desvio_texto.startswith("0.020235864105"):
        raise SystemExit(f"Desvio padrao inesperado: {desvio_texto}")


def gravar_resultados(resultado):
    linhas = [
        "medida;valor",
        f"quantidade;{resultado['quantidade']}",
        f"soma_reais;{resultado['soma']}",
        f"media;{resultado['media']}",
        f"mediana;{resultado['mediana']}",
        f"moda;{resultado['moda']}",
        f"frequencia_da_moda;{resultado['frequencia_moda']}",
        f"minimo;{resultado['minimo']}",
        f"maximo;{resultado['maximo']}",
        f"amplitude;{resultado['amplitude']}",
        f"quartil_1;{resultado['quartil_1']}",
        f"quartil_3;{resultado['quartil_3']}",
        f"intervalo_interquartil;{resultado['intervalo']}",
        f"limite_inferior;{resultado['limite_inferior']}",
        f"limite_superior;{resultado['limite_superior']}",
        f"soma_desvios_ao_quadrado;{resultado['soma_quadrados']}",
        f"variancia_amostral;{resultado['variancia']}",
        f"desvio_padrao_amostral;{resultado['desvio']}",
    ]
    destino = PASTA_ENTREGA / "resultados.csv"
    destino.write_text("\n".join(linhas) + "\n", encoding="utf-8")


def _borda():
    traco = Side(style="thin", color=LINHA)
    return Border(left=traco, right=traco, top=traco, bottom=traco)


def _cabecalho(celula):
    celula.font = Font(name="Calibri", bold=True, color="FFFFFF", size=12)
    celula.fill = PatternFill("solid", fgColor=AZUL)
    celula.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    celula.border = _borda()


def _texto(celula, negrito=False):
    celula.font = Font(name="Calibri", size=12, bold=negrito, color="1A1A1A")
    celula.alignment = Alignment(vertical="center", wrap_text=True)
    celula.border = _borda()


def gravar_planilha(tabela, resultado):
    livro = Workbook()
    dados = livro.active
    dados.title = "Dados"
    colunas = list(tabela.columns)
    for coluna, nome in enumerate(colunas, start=1):
        celula = dados.cell(1, coluna, nome)
        _cabecalho(celula)
    for linha, registro in enumerate(tabela.itertuples(index=False), start=2):
        for coluna, valor in enumerate(registro, start=1):
            celula = dados.cell(linha, coluna)
            if colunas[coluna - 1] == "preco_venda_rs_litro":
                celula.value = float(str(valor).replace(",", "."))
                celula.number_format = "0.00"
            else:
                celula.value = valor
            _texto(celula)
    dados.auto_filter.ref = "A1:J52"
    dados.freeze_panes = "A2"
    larguras = [16, 28, 16, 24, 14, 20, 55, 22, 14, 18]
    for indice, largura in enumerate(larguras, start=1):
        dados.column_dimensions[get_column_letter(indice)].width = largura
    dados.row_dimensions[1].height = 22

    medidas = livro.create_sheet("Medidas")
    medidas["A1"] = "Medida"
    medidas["B1"] = "Resultado da formula"
    medidas["C1"] = "Como a planilha calcula"
    for celula in (medidas["A1"], medidas["B1"], medidas["C1"]):
        _cabecalho(celula)

    formulas = [
        ("Quantidade de postos", "=COUNT(Dados!D2:D52)", "Conta os 51 precos da coluna D."),
        ("Soma dos precos (R$)", "=SUM(Dados!D2:D52)", "Soma os 51 precos. A media usa esta soma."),
        ("Media (R$/litro)", "=AVERAGE(Dados!D2:D52)", "Soma dividida por 51."),
        ("Mediana (R$/litro)", "=MEDIAN(Dados!D2:D52)", "Com 51 precos, e o 26o valor da lista ordenada."),
        ("Moda (R$/litro)", "=MODE(Dados!D2:D52)", "Preco que mais se repete."),
        ("Menor preco (R$/litro)", "=MIN(Dados!D2:D52)", "Menor valor do conjunto."),
        ("Maior preco (R$/litro)", "=MAX(Dados!D2:D52)", "Maior valor do conjunto."),
        ("Amplitude (R$/litro)", "=B8-B7", "Maior preco menos o menor preco."),
        ("Desvio padrao amostral", "=STDEV(Dados!D2:D52)", "Raiz da variancia. O divisor e 51 - 1 = 50."),
        ("Variancia amostral", "=VAR(Dados!D2:D52)", "Soma dos desvios ao quadrado, dividida por 50."),
        ("Q1, metodo da aula", "=SMALL(Dados!D2:D52,13)", "13o preco da lista ordenada. Centro da metade de baixo."),
        ("Q3, metodo da aula", "=SMALL(Dados!D2:D52,39)", "39o preco da lista ordenada. Centro da metade de cima."),
        ("Intervalo interquartil", "=B13-B12", "Q3 menos Q1."),
        ("Limite inferior do boxplot", "=B12-1.5*B14", "Q1 menos 1,5 vezes o intervalo interquartil."),
        ("Limite superior do boxplot", "=B13+1.5*B14", "Q3 mais 1,5 vezes o intervalo interquartil."),
        ("Soma dos desvios ao quadrado", "=DEVSQ(Dados!D2:D52)", "Base manual do desvio padrao."),
        ("Desvio conferido na mao", "=SQRT(B17/(B2-1))", "Deve bater com o desvio padrao amostral."),
    ]
    for linha, (nome, formula, explicacao) in enumerate(formulas, start=2):
        medidas.cell(linha, 1, nome)
        medidas.cell(linha, 2, formula)
        medidas.cell(linha, 3, explicacao)
        for coluna in range(1, 4):
            _texto(medidas.cell(linha, coluna), negrito=(coluna == 1))
        medidas.cell(linha, 2).alignment = Alignment(horizontal="center", vertical="center")
    for linha in (3, 5, 6, 7, 8, 9, 12, 13):
        medidas.cell(linha, 2).number_format = "0.00"
    for linha in (4, 10, 14, 15, 16, 18):
        medidas.cell(linha, 2).number_format = "0.0000"
    medidas.cell(11, 2).number_format = "0.000000"
    medidas.cell(17, 2).number_format = "0.00000000"
    medidas.cell(2, 2).number_format = "0"
    medidas.column_dimensions["A"].width = 36
    medidas.column_dimensions["B"].width = 28
    medidas.column_dimensions["C"].width = 68
    medidas.row_dimensions[1].height = 22
    for linha in range(2, 19):
        medidas.row_dimensions[linha].height = 20
        if linha % 2 == 0:
            for coluna in range(1, 4):
                medidas.cell(linha, coluna).fill = PatternFill("solid", fgColor=CINZA)

    frequencia = livro.create_sheet("Frequencia")
    for coluna, nome in enumerate(("Preco (R$/litro)", "Quantidade de postos", "Percentual"), start=1):
        _cabecalho(frequencia.cell(1, coluna, nome))
    precos_unicos = sorted(resultado["frequencia"])
    for indice, preco in enumerate(precos_unicos, start=2):
        frequencia.cell(indice, 1, float(preco))
        frequencia.cell(indice, 2, f'=COUNTIF(Dados!D$2:D$52,A{indice})')
        frequencia.cell(indice, 3, f"=B{indice}/Medidas!B2")
        frequencia.cell(indice, 1).number_format = "0.00"
        frequencia.cell(indice, 3).number_format = "0.00%"
        for coluna in range(1, 4):
            _texto(frequencia.cell(indice, coluna))
            frequencia.cell(indice, coluna).alignment = Alignment(horizontal="center")
    ultima = 1 + len(precos_unicos)
    frequencia.cell(ultima + 1, 1, "Total")
    frequencia.cell(ultima + 1, 2, f"=SUM(B2:B{ultima})")
    frequencia.cell(ultima + 1, 3, f"=SUM(C2:C{ultima})")
    frequencia.cell(ultima + 1, 3).number_format = "0.00%"
    for coluna in range(1, 4):
        _texto(frequencia.cell(ultima + 1, coluna), negrito=True)
    frequencia.column_dimensions["A"].width = 24
    frequencia.column_dimensions["B"].width = 26
    frequencia.column_dimensions["C"].width = 18

    grafico = BarChart()
    grafico.type = "col"
    grafico.title = "Quantos postos cobraram cada preco"
    grafico.y_axis.title = "Quantidade de postos"
    grafico.x_axis.title = "Preco da gasolina comum (R$/litro)"
    dados_grafico = Reference(frequencia, min_col=2, min_row=1, max_row=ultima)
    rotulos = Reference(frequencia, min_col=1, min_row=2, max_row=ultima)
    grafico.add_data(dados_grafico, titles_from_data=True)
    grafico.set_categories(rotulos)
    grafico.shape = 4
    grafico.legend = None
    grafico.y_axis.scaling.min = 0
    grafico.y_axis.scaling.max = 24
    grafico.y_axis.majorUnit = 3
    grafico.dataLabels = DataLabelList()
    grafico.dataLabels.showVal = True
    grafico.style = 10
    grafico.width = 18
    grafico.height = 8
    frequencia.add_chart(grafico, "E2")

    ordem = livro.create_sheet("Lista_ordenada")
    _cabecalho(ordem.cell(1, 1, "Posicao"))
    _cabecalho(ordem.cell(1, 2, "Preco ordenado (R$/litro)"))
    _cabecalho(ordem.cell(1, 3, "Papel no calculo"))
    papeis = {
        13: "Q1, centro da metade de baixo",
        26: "Mediana, 26a posicao",
        39: "Q3, centro da metade de cima",
    }
    for posicao in range(1, 52):
        ordem.cell(posicao + 1, 1, posicao)
        ordem.cell(posicao + 1, 2, f"=SMALL(Dados!$D$2:$D$52,A{posicao + 1})")
        ordem.cell(posicao + 1, 3, papeis.get(posicao, ""))
        ordem.cell(posicao + 1, 2).number_format = "0.00"
        for coluna in range(1, 4):
            _texto(ordem.cell(posicao + 1, coluna))
        if posicao in papeis:
            for coluna in range(1, 4):
                ordem.cell(posicao + 1, coluna).fill = PatternFill("solid", fgColor="F4E4C4")
    ordem.column_dimensions["A"].width = 14
    ordem.column_dimensions["B"].width = 32
    ordem.column_dimensions["C"].width = 40
    ordem.freeze_panes = "A2"

    notas = livro.create_sheet("Como_conferir")
    notas["A1"] = "Conferencia dos calculos"
    notas["A1"].font = Font(name="Calibri", bold=True, size=16, color=AZUL)
    textos = [
        "Abra esta planilha no LibreOffice ou no Excel e deixe as formulas calcularem.",
        "A coluna B da aba Medidas deve mostrar: quantidade 51, soma 355,21, media 6,9649, mediana 6,97, moda 6,95, amplitude 0,10 e desvio padrao amostral 0,0202.",
        "O desvio usa a versao amostral. O divisor e n - 1, isto e, 50. A celula B18 refaz essa conta com a soma dos desvios ao quadrado.",
        "Quartis: n = 51 e impar. A mediana e o 26o preco. Ela sai da lista antes de achar Q1 e Q3. Cada metade fica com 25 precos. Q1 e Q3 sao o 13o preco de cada metade, nas posicoes 13 e 39 da lista completa.",
        "Esse metodo e o da aula. Outro metodo de quartil pode dar numero diferente.",
        "Nenhum preco foi tirado da base. O preco de 6,89 fica no limite inferior do boxplot e continua no conjunto.",
        "Fonte dos precos: ANP, serie historica de precos de combustiveis, setembro de 2026. Recorte de 21 a 27/09/2026. Consulta em 08/10/2026.",
    ]
    for linha, texto in enumerate(textos, start=3):
        notas.cell(linha, 1, texto)
        notas.cell(linha, 1).font = Font(name="Calibri", size=12)
        notas.cell(linha, 1).alignment = Alignment(wrap_text=True, vertical="center")
        notas.row_dimensions[linha].height = 36
    notas.column_dimensions["A"].width = 120
    notas.row_dimensions[1].height = 24
    notas.sheet_properties.tabColor = AZUL

    medidas.sheet_properties.tabColor = AZUL
    livro.save(PASTA_ENTREGA / "conferencia_calculos.xlsx")


def _eixo_reais(eixo):
    eixo.set_major_formatter(FuncFormatter(lambda valor, _posicao: f"{valor:.2f}".replace(".", ",")))


def gravar_graficos(resultado):
    plt.rcParams["font.family"] = "DejaVu Sans"
    ordem = [float(preco) for preco in resultado["ordem"]]
    faixas = [6.88, 6.90, 6.92, 6.94, 6.96, 6.98, 7.00]
    rotulos = [
        "6,88 a\nmenos de 6,90",
        "6,90 a\nmenos de 6,92",
        "6,92 a\nmenos de 6,94",
        "6,94 a\nmenos de 6,96",
        "6,96 a\nmenos de 6,98",
        "6,98 a\nmenos de 7,00",
    ]
    contagens = [0] * (len(faixas) - 1)
    for preco in ordem:
        for indice in range(len(faixas) - 1):
            if faixas[indice] <= preco < faixas[indice + 1]:
                contagens[indice] += 1
                break
    if sum(contagens) != 51:
        raise SystemExit(f"Histograma nao fechou em 51: {contagens}")

    figura, eixo = plt.subplots(figsize=(10, 5.8))
    barras = eixo.bar(range(len(contagens)), contagens, width=0.92, color="#2F6FAD", edgecolor="white")
    eixo.set_xticks(range(len(rotulos)))
    eixo.set_xticklabels(rotulos)
    eixo.set_ylabel("Quantidade de postos")
    eixo.set_xlabel("Preço da gasolina comum (R$/litro)")
    eixo.set_title("Onde os preços da gasolina comum se concentram")
    eixo.set_ylim(0, 26)
    eixo.yaxis.grid(True, linestyle=":", color="#C5D0DC")
    eixo.set_axisbelow(True)
    for barra, contagem in zip(barras, contagens):
        eixo.text(barra.get_x() + barra.get_width() / 2, contagem + 0.4, str(contagem), ha="center", va="bottom", fontsize=11)
    eixo.spines["top"].set_visible(False)
    eixo.spines["right"].set_visible(False)
    figura.tight_layout()
    figura.savefig(PASTA_ENTREGA / "histograma.png", dpi=160, bbox_inches="tight")
    plt.close(figura)

    stats = [
        {
            "med": float(resultado["mediana"]),
            "q1": float(resultado["quartil_1"]),
            "q3": float(resultado["quartil_3"]),
            "whislo": float(resultado["minimo"]),
            "whishi": float(resultado["maximo"]),
            "fliers": [],
            "mean": float(resultado["media"]),
        }
    ]
    figura, eixo = plt.subplots(figsize=(10, 4.2))
    eixo.bxp(
        stats,
        orientation="horizontal",
        showmeans=True,
        meanprops={"marker": "D", "markerfacecolor": "#C47B17", "markeredgecolor": "#C47B17", "markersize": 7},
        boxprops={"facecolor": "#D6E6F5", "edgecolor": "#1F4E79", "linewidth": 1.4},
        medianprops={"color": "#B23A48", "linewidth": 2},
        whiskerprops={"color": "#1F4E79", "linewidth": 1.3},
        capprops={"color": "#1F4E79", "linewidth": 1.3},
        patch_artist=True,
        widths=0.45,
    )
    eixo.set_yticks([1])
    eixo.set_yticklabels(["51 postos"])
    eixo.set_xlabel(
        "Preço da gasolina comum (R$/litro)\n"
        "Caixa: de R$ 6,95 a R$ 6,99. Linha vermelha: mediana R$ 6,97. Losango: média R$ 6,96.",
        fontsize=11,
    )
    eixo.set_title("Boxplot do preço da gasolina comum")
    eixo.set_xlim(6.86, 7.02)
    _eixo_reais(eixo.xaxis)
    eixo.xaxis.grid(True, linestyle=":", color="#C5D0DC")
    eixo.set_axisbelow(True)
    eixo.spines["top"].set_visible(False)
    eixo.spines["right"].set_visible(False)
    eixo.annotate(
        "menor preço\nR$ 6,89",
        xy=(float(resultado["minimo"]), 1),
        xytext=(6.89, 1.48),
        ha="center",
        va="bottom",
        fontsize=10,
        arrowprops={"arrowstyle": "-", "color": "#5C6B7A"},
        color="#1A1A1A",
    )
    eixo.set_ylim(0.55, 1.85)
    figura.tight_layout()
    figura.savefig(PASTA_ENTREGA / "boxplot.png", dpi=160, bbox_inches="tight")
    plt.close(figura)


def main():
    PASTA_ENTREGA.mkdir(parents=True, exist_ok=True)
    tabela, precos = ler_precos()
    if tabela["produto"].nunique() != 1 or tabela["cnpj_revenda"].duplicated().any():
        raise SystemExit("A base nao passou na conferencia de produto unico e CNPJ unico.")
    resultado = calcular(precos)
    conferir(resultado)
    gravar_resultados(resultado)
    gravar_planilha(tabela, resultado)
    gravar_graficos(resultado)
    print("quantidade", resultado["quantidade"])
    print("media", dinheiro(resultado["media"], "0.0001"))
    print("mediana", resultado["mediana"])
    print("moda", resultado["moda"], resultado["frequencia_moda"])
    print("amplitude", resultado["amplitude"])
    print("desvio", dinheiro(resultado["desvio"], "0.0001"))
    print("q1", resultado["quartil_1"], "q3", resultado["quartil_3"])


if __name__ == "__main__":
    main()
